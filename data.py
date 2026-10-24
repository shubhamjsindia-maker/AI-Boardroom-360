"""Workbook template and standardized import. INR lakh. No macros executed."""
from io import BytesIO
from dataclasses import asdict
import pandas as pd
from openpyxl import Workbook
from finance import Inputs, validate

SHEETS = {
 'Company_Master': ['company_id','name','industry','financial_year_end','currency','unit'],
 'Historical_Financials': ['year','revenue','direct_cost','operating_expenses','ebitda','depreciation','interest','tax','profit'],
 'Balance_Sheet': ['year','cash','receivables','inventory','fixed_assets','other_assets','equity','borrowings','payables','other_liabilities'],
 'Existing_Debt': ['loan_id','principal_outstanding','interest_rate','remaining_years','principal_per_year'],
 'Working_Capital': ['year','receivable_days','inventory_days','payable_days'],
 'Scenario_Inputs': ['field','value']
}

NAME='Nova Manufacturing Pvt. Ltd.'
def sample_inputs(): return Inputs()

def template_bytes():
    i=sample_inputs(); w=Workbook(); w.remove(w.active)
    rows={
      'Company_Master':[['NOVA001',NAME,'Manufacturing','2026-03-31','INR','lakh']],
      'Historical_Financials':[[2024,830,456.5,257.5,116,16,15,21.25,63.75],[2025,920,506,276,138,18,13,26.75,80.25],[2026,1000,550,300,150,20,10,30,90]],
      'Balance_Sheet':[[2026,80,123.29,52.74,400,43.97,350,100,45.21,204.79]],
      'Existing_Debt':[['L001',100,0.10,5,20]],
      'Working_Capital':[[2026,45,35,30]],
      'Scenario_Inputs':[[k,v] for k,v in asdict(i).items()]
    }
    # balance sheet sample balanced: 80+123.29+52.74+400+43.97=700; RHS 350+100+45.21+204.79=700
    for s,headers in SHEETS.items():
        ws=w.create_sheet(s);ws.append(headers)
        for r in rows[s]: ws.append(r)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]: cell.font=__import__('openpyxl').styles.Font(bold=True,color='FFFFFF');cell.fill=__import__('openpyxl').styles.PatternFill('solid',fgColor='16324B')
        for col in ws.columns:
            letter=col[0].column_letter
            ws.column_dimensions[letter].width=min(42,max(17,max(len(str(c.value or '')) for c in col)+3))
    bio=BytesIO();w.save(bio);return bio.getvalue()


def parse_workbook(blob):
    if len(blob)>5*1024*1024: raise ValueError('Maximum workbook size is 5 MB')
    if not blob[:2]==b'PK': raise ValueError('Only .xlsx files supported')
    book=pd.ExcelFile(BytesIO(blob),engine='openpyxl')
    problems=[]
    for s,fields in SHEETS.items():
        if s not in book.sheet_names: problems.append(f'Missing sheet: {s}');continue
        df=pd.read_excel(book,sheet_name=s)
        missing=set(fields)-set(df.columns)
        if missing: problems.append(f'{s}: missing columns {sorted(missing)}')
    if problems: raise ValueError('; '.join(problems))
    company=pd.read_excel(book,sheet_name='Company_Master').iloc[0].to_dict()
    if str(company['currency']).upper()!='INR' or str(company['unit']).lower()!='lakh':
        raise ValueError('MVP accepts INR lakh only; convert source units explicitly')
    historical=pd.read_excel(book,sheet_name='Historical_Financials')
    bs=pd.read_excel(book,sheet_name='Balance_Sheet')
    debt=pd.read_excel(book,sheet_name='Existing_Debt')
    wc=pd.read_excel(book,sheet_name='Working_Capital')
    assumptions=pd.read_excel(book,sheet_name='Scenario_Inputs')
    for name,frame in [('Historical_Financials',historical),('Balance_Sheet',bs),('Working_Capital',wc)]:
        if frame['year'].duplicated().any(): raise ValueError(f'{name}: duplicate year')
    if historical.empty or bs.empty or debt.empty or wc.empty: raise ValueError('Financial tables cannot be empty')
    h=historical.sort_values('year').iloc[-1];b=bs.sort_values('year').iloc[-1]; d=debt.iloc[0]; working=wc.sort_values('year').iloc[-1]
    if int(h['year'])!=int(b['year']) or int(h['year'])!=int(working['year']):raise ValueError('Latest reporting years must match')
    assets=sum(float(b[k]) for k in ['cash','receivables','inventory','fixed_assets','other_assets'])
    liabilities=sum(float(b[k]) for k in ['equity','borrowings','payables','other_liabilities'])
    if abs(assets-liabilities)>0.1:raise ValueError(f'Balance sheet does not reconcile: difference ₹{assets-liabilities:.2f} lakh')
    if abs(float(h['revenue'])-float(h['direct_cost'])-float(h['operating_expenses'])-float(h['ebitda']))>0.1:
        raise ValueError('EBITDA does not reconcile with historic revenue and costs')
    if abs(float(d['principal_outstanding'])-float(b['borrowings']))>0.1:raise ValueError('Existing loan differs from balance-sheet borrowing')
    a=dict(zip(assumptions.field,assumptions.value))
    if set(a)!=set(asdict(Inputs())): raise ValueError('Scenario_Inputs must contain exactly the defined input names')
    vals={k:float(v) for k,v in a.items()}
    vals.update(revenue=float(h['revenue']), baseline_ebitda_margin=float(h['ebitda'])/float(h['revenue']),
                variable_cost_ratio=float(h['direct_cost'])/float(h['revenue']),cash=float(b['cash']),debt=float(b['borrowings']),
                interest_existing=float(d['interest_rate']),existing_principal_per_year=float(d['principal_per_year']),
                receivable_days=float(working['receivable_days']),inventory_days=float(working['inventory_days']),payable_days=float(working['payable_days']))
    vals['term_years']=int(vals['term_years']);vals['depreciation_years']=int(vals['depreciation_years']);vals['capex_start_year']=int(vals['capex_start_year'])
    i=Inputs(**vals)
    if validate(i): raise ValueError('; '.join(validate(i)))
    warnings=[]
    if len(historical)<3: warnings.append('Fewer than 3 years of historical financial data')
    if float(b['cash'])<0: warnings.append('Negative cash balance')
    return str(company['name']),i,warnings
